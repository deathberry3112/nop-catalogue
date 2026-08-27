#!/usr/bin/env python3
"""Phase 2's sheet — locations, areas, and what the parser refuses to read.

One sheet line is one AREA. Lines sharing a SKU are one product; lines sharing a
(SKU, Location Name) are one location, and only the first of them carries the
Picario reference — the blank cell on the lines after it is what joins them to
the location above rather than starting another one. Most of what is worth
testing is that assembly, and that a sheet which cannot mean what it appears to
mean is refused rather than read.

The round-trip tests matter for the same reason they do in test_parser: this
module and xlsxCreator are mirrors, and the report is meant to be re-uploaded to
finish a batch.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

import xlsxCreator
from app import configurator
from app.errors import ErrorClass, classify
from app.parser import parse_workbook

HEADERS = xlsxCreator.CONFIGURATOR_HEADERS


def _workbook(rows: list[list], headers: list[str] | None = None,
              sheet: str = "Configurator Setup") -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers or HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _line(sku="G347LA", location="Front", ref="", picario="A-front",
          thumb="A-front.png", *, name=None, save_product="Yes") -> list:
    """One area line. `ref` empty means "another area on the location above"."""
    return [name if name is not None else sku, sku, "Select",
            "Yes" if ref else "", location, ref, "Yes" if ref else "",
            picario, thumb, "Yes", save_product]


def _rows(source) -> list[dict]:
    return [r for r in parse_workbook(source).rows
            if r.get("kind") == "configurator"]


def _one(source) -> dict:
    rows = _rows(source)
    assert len(rows) == 1, [r["sku"] for r in rows]
    return rows[0]


# --- telling the two templates apart ----------------------------------------

def test_a_sheet_with_location_and_reference_columns_is_a_configurator_sheet():
    assert configurator.looks_like_configurator(tuple(HEADERS))


def test_the_product_template_is_not_mistaken_for_a_configurator_sheet():
    assert not configurator.looks_like_configurator(
        tuple(xlsxCreator.TEMPLATE_HEADERS))


def test_detection_is_by_header_not_by_sheet_name():
    """The client calls the tab "Configurator Setup"; the columns are what
    decide, so the same sheet under any other name still reads."""
    row = _one(_workbook([_line(ref="G347LA-Front.pfs")], sheet="Anything At All"))
    assert row["locations"][0]["reference_id"] == "G347LA-Front.pfs"


def test_the_two_templates_can_share_one_workbook():
    wb = openpyxl.Workbook()
    products = wb.active
    products.title = "Products"
    products.append(xlsxCreator.TEMPLATE_HEADERS)
    products.append(["P", "SKU-1", "Zing", "Color", "Red", "Red", "Size", "",
                     "Yes", "No", "SKU-1 RD", "Yes"])
    areas = wb.create_sheet("Configurator Setup")
    areas.append(HEADERS)
    areas.append(_line(sku="SKU-1", ref="a.pfs"))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    result = parse_workbook(buf)
    kinds = sorted((r.get("kind") or "product") for r in result.rows)
    assert kinds == ["configurator", "product"]


def test_a_product_row_and_a_configurator_row_may_share_a_sku():
    """Phase 1 and phase 2 of the same product. Counting them as one duplicated
    SKU would fail every product the client configures."""
    wb = openpyxl.Workbook()
    products = wb.active
    products.title = "Products"
    products.append(xlsxCreator.TEMPLATE_HEADERS)
    products.append(["P", "G55FF", "Zing", "Color", "Red", "Red", "Size", "",
                     "Yes", "No", "G55FF RD", "Yes"])
    areas = wb.create_sheet("Configurator Setup")
    areas.append(HEADERS)
    areas.append(_line(sku="G55FF", ref="G55FF-Front.pfs"))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    for row in parse_workbook(buf).rows:
        assert row["errors"] == [], row["errors"]


# --- assembling lines into locations and areas ------------------------------

def test_lines_sharing_a_location_become_one_location():
    row = _one(_workbook([
        _line(location="Front", ref="G347LA-Front.pfs", picario="L-chest",
              thumb="L-chest.png"),
        _line(location="Front", picario="K-chest", thumb="K-chest.png"),
    ]))
    assert len(row["locations"]) == 1
    location = row["locations"][0]
    assert location["reference_id"] == "G347LA-Front.pfs"
    assert [a["picario_name"] for a in location["areas"]] == ["L-chest", "K-chest"]


def test_a_new_reference_starts_a_new_location():
    row = _one(_workbook([
        _line(location="Front", ref="F.pfs"),
        _line(location="Front"),
        _line(location="Back", ref="B.pfs"),
        _line(location="Back"),
    ]))
    assert [(l["name"], l["reference_id"]) for l in row["locations"]] == [
        ("Front", "F.pfs"), ("Back", "B.pfs")]
    assert [len(l["areas"]) for l in row["locations"]] == [2, 2]


def test_lines_are_grouped_by_sku_not_by_adjacency():
    """A sheet sorted by location rather than by product still assembles."""
    rows = _rows(_workbook([
        _line(sku="A", location="Front", ref="A-F.pfs"),
        _line(sku="B", location="Front", ref="B-F.pfs"),
        _line(sku="A", location="Back", ref="A-B.pfs"),
        _line(sku="B", location="Back", ref="B-B.pfs"),
    ]))
    assert {r["sku"] for r in rows} == {"A", "B"}
    for row in rows:
        assert [l["name"] for l in row["locations"]] == ["Front", "Back"]


def test_area_order_within_a_location_is_the_sheets_order():
    """The plugin stores a DisplayOrder per area, so the sheet's sequence is the
    answer to which comes first."""
    row = _one(_workbook([
        _line(ref="F.pfs", picario="one", thumb="one.png"),
        _line(picario="two", thumb="two.png"),
        _line(picario="three", thumb="three.png"),
    ]))
    assert [a["picario_name"] for a in row["locations"][0]["areas"]] == [
        "one", "two", "three"]


def test_every_source_row_is_recorded():
    row = _one(_workbook([_line(ref="F.pfs"), _line(), _line()]))
    assert row["source_rows"] == [2, 3, 4]
    assert [a["source_row"] for a in row["locations"][0]["areas"]] == [2, 3, 4]


# --- the cells themselves ---------------------------------------------------

def test_internal_spacing_in_a_reference_is_preserved():
    """"280-HP542R F WE .pfs" is a real Picario scene name, space and all.
    Tidying it would ask Picario for a scene that does not exist."""
    row = _one(_workbook([_line(ref="280-HP542R F WE .pfs")]))
    assert row["locations"][0]["reference_id"] == "280-HP542R F WE .pfs"


def test_a_thumbnail_is_prefixed_the_way_the_admin_lists_it():
    row = _one(_workbook([_line(ref="F.pfs", thumb="A-front-hard-hat.png")]))
    area = row["locations"][0]["areas"][0]
    assert area["thumbnail"] == "/thumbnails/A-front-hard-hat.png"
    assert area["thumbnail_asked"] == "A-front-hard-hat.png"


def test_an_already_prefixed_thumbnail_is_not_prefixed_twice():
    row = _one(_workbook([_line(ref="F.pfs", thumb="/thumbnails/A.png")]))
    assert row["locations"][0]["areas"][0]["thumbnail"] == "/thumbnails/A.png"


def test_the_navigation_columns_are_not_treated_as_data():
    """'Configurator Areas: Select' and 'Add new location: Yes' tell a human
    where to click. They must not trip the unknown-column guard."""
    row = _one(_workbook([_line(ref="F.pfs")]))
    assert row["errors"] == []


# --- what it refuses --------------------------------------------------------

def test_a_location_with_no_picario_reference_is_refused():
    row = _one(_workbook([_line(location="Front", picario="a", thumb="a.png")]))
    assert any("no Picario RefrenceId" in e for e in row["errors"])


def test_a_location_given_two_different_references_is_refused():
    """Nothing here knows which of the two answers to discard."""
    row = _one(_workbook([
        _line(location="Front", ref="one.pfs"),
        _line(location="Front", ref="two.pfs"),
    ]))
    assert any("two different Picario references" in e.replace("2 different", "two different")
               or "different Picario references" in e for e in row["errors"])


def test_the_same_reference_repeated_on_every_line_is_fine():
    """Repetition is not contradiction — it re-parses as the same location."""
    row = _one(_workbook([
        _line(location="Front", ref="F.pfs", picario="a", thumb="a.png"),
        _line(location="Front", ref="F.pfs", picario="b", thumb="b.png"),
    ]))
    assert row["errors"] == []
    assert len(row["locations"]) == 1
    assert len(row["locations"][0]["areas"]) == 2


def test_two_areas_on_one_spot_are_refused():
    """The plugin refuses this itself, but only after the location exists."""
    row = _one(_workbook([
        _line(ref="F.pfs", picario="A-front", thumb="one.png"),
        _line(picario="A-front", thumb="two.png"),
    ]))
    assert any("two areas on the same spot" in e for e in row["errors"])


def test_the_same_spot_on_two_different_locations_is_fine():
    """Front and Back are different scenes; a name is only unique within one."""
    row = _one(_workbook([
        _line(location="Front", ref="F.pfs", picario="A-front"),
        _line(location="Back", ref="B.pfs", picario="A-front"),
    ]))
    assert row["errors"] == []


def test_an_area_with_no_picario_name_is_refused():
    row = _one(_workbook([_line(ref="F.pfs", picario="")]))
    assert any("no Picario Name" in e for e in row["errors"])


def test_an_area_with_no_thumbnail_is_refused():
    row = _one(_workbook([_line(ref="F.pfs", thumb="")]))
    assert any("no Thumbnail Name" in e for e in row["errors"])


def test_a_line_with_no_sku_is_refused_on_its_own():
    rows = _rows(_workbook([
        _line(sku="A", ref="A.pfs"),
        _line(sku="", ref="B.pfs"),
    ]))
    blank = [r for r in rows if not r["sku"]]
    assert len(blank) == 1
    assert "SKU is required" in blank[0]["errors"]
    assert [r for r in rows if r["sku"] == "A"][0]["errors"] == []


def test_a_column_nobody_models_refuses_the_row():
    """The same silent-drop guard the product sheet applies: a field this
    pipeline cannot set produces a wrong storefront and a clean report."""
    row = _one(_workbook([_line(ref="F.pfs") + ["something"]],
                         headers=HEADERS + ["Mystery Column"]))
    assert any("unrecognised column" in e for e in row["errors"])


def test_a_sheet_with_no_sku_column_is_refused_whole():
    headers = [h for h in HEADERS if h != "SKU"]
    result = parse_workbook(_workbook([["G347LA", "Select", "Yes", "Front",
                                        "F.pfs", "Yes", "a", "a.png", "Yes",
                                        "Yes"]], headers=headers))
    assert any("no 'SKU' column" in e for e in result.errors)
    assert result.rows == []


# --- the report, and re-uploading it ----------------------------------------

def test_the_report_round_trips_a_configurator_row(tmp_path):
    source = _workbook([
        _line(location="Front", ref="F.pfs", picario="one", thumb="one.png"),
        _line(location="Front", picario="two", thumb="two.png"),
        _line(location="Back", ref="B.pfs", picario="three", thumb="three.png"),
    ])
    before = _one(source)

    out = tmp_path / "report.xlsx"
    xlsxCreator.write_report([before], None, out)
    after = _one(out)

    def shape(row):
        return [(l["name"], l["reference_id"],
                 [(a["picario_name"], a["thumbnail"]) for a in l["areas"]])
                for l in row["locations"]]

    assert shape(after) == shape(before)


def test_the_report_writes_the_reference_only_on_a_locations_first_line(tmp_path):
    """Writing it on none of them would not re-parse; that is what this guards."""
    row = _one(_workbook([
        _line(location="Front", ref="F.pfs", picario="one"),
        _line(location="Front", picario="two"),
    ]))
    lines = xlsxCreator.configurator_to_lines(row)
    reference_column = xlsxCreator.CONFIGURATOR_HEADERS.index("Picario RefrenceId")
    assert [line[reference_column] for line in lines] == ["F.pfs", ""]


def test_a_configured_row_is_not_run_again_when_the_report_is_re_uploaded(tmp_path):
    import main

    row = _one(_workbook([_line(sku="G55FF", ref="F.pfs")]))
    out = tmp_path / "report.xlsx"
    summary = {"results": [{"sku": "G55FF", "kind": "configurator",
                            "status": "Configured", "error": None}]}
    xlsxCreator.write_report([row], summary, out)

    runnable, done, blocked = main.partition(_rows(out))
    assert [r["sku"] for r in done] == ["G55FF"]
    assert runnable == [] and blocked == []


def test_a_products_status_never_lands_on_its_configurator_row(tmp_path):
    """Both rows may name one SKU, so the report matches on (kind, SKU)."""
    product = {"sku": "G55FF", "source_sheet": "Products", "colors": ["Red"],
               "sizes": [], "combinations": [], "variants": [], "errors": []}
    areas = _one(_workbook([_line(sku="G55FF", ref="F.pfs")]))
    summary = {"results": [
        {"sku": "G55FF", "kind": "product", "error": "the form rejected this row"},
        {"sku": "G55FF", "kind": "configurator", "status": "Configured",
         "error": None},
    ]}
    index = {xlsxCreator._result_key(r): r for r in summary["results"]}

    assert xlsxCreator.status_for(product, index)[0] == xlsxCreator.STATUS_FAILED
    assert xlsxCreator.status_for(areas, index)[0] == "Configured"


# --- failure classification --------------------------------------------------

@pytest.mark.parametrize("message", [
    "no product with SKU 'G347LA' in this catalogue",
    "configurator: the admin would not accept Picario reference 'x.pfs'",
    "configurator: Picario scene for this location has no area named 'x'",
    "configurator: this store has no thumbnail '/thumbnails/x.png'",
])
def test_phase_two_data_failures_are_deterministic_and_never_retried(message):
    """Five attempts would spend five minutes reaching the same answer: the
    reference, the spot and the thumbnail are facts about Picario and the
    catalogue, not about the browser."""
    classified = classify(message)
    assert classified.error_class is ErrorClass.DATA_INVALID
    assert not classified.retryable


# --- alternative Picario references -----------------------------------------
# A reference the workbook writes may name a scene that does not contain the
# spot the workbook then asks for. `reference_alternatives` proposes the one
# other spelling worth trying; ConfiguratorPage only ever ACCEPTS it after
# seeing that the alternative scene really does hold that spot.

def test_one_space_before_the_extension_is_the_candidate():
    assert configurator.reference_alternatives("280-HP542R F WE .pfs") == [
        "280-HP542R F WE.pfs"]


def test_a_reference_with_no_space_before_the_extension_has_no_alternative():
    """Nothing to try, so nothing is tried — the written reference stands."""
    assert configurator.reference_alternatives("280-HP542R B WE.pfs") == []
    assert configurator.reference_alternatives("G55FF-Front.pfs") == []


def test_internal_spaces_elsewhere_are_left_alone():
    """Only the space against the extension is a candidate. "280-HP542R F WE"
    is the scene's real name and collapsing it would invent a filename."""
    assert configurator.reference_alternatives("280-HP542R F WE .pfs") == [
        "280-HP542R F WE.pfs"]
    assert configurator.reference_alternatives("a b c.pfs") == []


def test_only_one_space_is_removed():
    """One space, one position — not a whitespace strip."""
    assert configurator.reference_alternatives("a  .pfs") == ["a .pfs"]


def test_a_reference_with_no_extension_has_no_alternative():
    assert configurator.reference_alternatives("no-extension ") == []
    assert configurator.reference_alternatives("") == []


def test_the_alternative_never_changes_what_the_row_carries():
    """The sheet's own value is what the row holds; the alternative is only ever
    proposed at run time, against a scene that can confirm it."""
    row = _one(_workbook([_line(ref="280-HP542R F WE .pfs")]))
    assert row["locations"][0]["reference_id"] == "280-HP542R F WE .pfs"
