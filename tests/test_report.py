#!/usr/bin/env python3
"""The .xlsx report's Status/Note columns. No browser, no credentials.

Pins the contract test_add_products relies on when it turns a run's outcome
into a summary for xlsxCreator.
"""

from __future__ import annotations

import openpyxl
import pytest

import main
import xlsxCreator


@pytest.fixture(scope="module")
def payload() -> list[dict]:
    return main.load_rows(main.DEFAULT_PAYLOAD)


def read_back(path):
    """{sku: (status, note)} from a written report, across every sheet.

    An empty Note comes back from openpyxl as None rather than "", so it is
    normalised here — the distinction is a spreadsheet artefact, not something
    the report is trying to say.
    """
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        sku_col = headers.index("SKU")
        status_col = headers.index("Status")
        note_col = headers.index("Note")
        for row in ws.iter_rows(min_row=2, values_only=True):
            out[row[sku_col]] = (row[status_col], row[note_col] or "")
    return out


def test_every_payload_row_appears(payload, tmp_path):
    path = tmp_path / "report.xlsx"
    xlsxCreator.write_report(payload, {"results": []}, path)
    assert set(read_back(path)) == {r["sku"] for r in payload}


def test_sheets_mirror_the_source_workbook(payload, tmp_path):
    path = tmp_path / "report.xlsx"
    xlsxCreator.write_report(payload, {"results": []}, path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["HardHat", "Vest"]


def test_the_four_outcomes(payload, tmp_path):
    """Added / Failed / Skipped / Pending, exactly as the add test emits them."""
    path = tmp_path / "report.xlsx"
    summary = {"results": [
        {"sku": "CP-30-C100"},                                  # created
        {"sku": "CP-30-F100", "status": xlsxCreator.STATUS_ADDED,
         "error": "created by an earlier run; configuration completed"},
        {"sku": "CP-30-C100V", "status": xlsxCreator.STATUS_SKIPPED,
         "error": "already in the catalogue (.../Edit/388) and not created by "
                  "this automation — left untouched"},
        {"sku": "CP-V-010", "error": "combinations do not line up"},
        # CP-30-F100V, CP-V-010T, CP-V-200 deliberately absent -> Pending
    ]}
    xlsxCreator.write_report(payload, summary, path)
    got = read_back(path)

    assert got["CP-30-C100"] == ("Added", "")
    assert got["CP-30-F100"][0] == "Added"
    assert "earlier run" in got["CP-30-F100"][1]
    assert got["CP-30-C100V"][0] == "Skipped"
    assert "left untouched" in got["CP-30-C100V"][1]
    assert got["CP-V-010"] == ("Failed", "combinations do not line up")

    for sku in ("CP-30-F100V", "CP-V-010T", "CP-V-200"):
        assert got[sku][0] == "Pending", f"{sku} was never reached"


def test_a_halted_run_still_reports_what_it_did(payload, tmp_path):
    """The case that matters: the run died after two products. Those two read
    Added, the rest read Pending — nobody has to guess where it stopped."""
    path = tmp_path / "report.xlsx"
    done = [r["sku"] for r in payload[:2]]
    xlsxCreator.write_report(payload, {"results": [{"sku": s} for s in done]}, path)
    got = read_back(path)

    assert [got[s][0] for s in done] == ["Added", "Added"]
    assert all(got[r["sku"]][0] == "Pending" for r in payload[2:])


def test_the_payload_columns_survive_the_round_trip(payload, tmp_path):
    """The report is re-uploadable, so the 12 template columns must come back
    out intact alongside Status/Note."""
    path = tmp_path / "report.xlsx"
    xlsxCreator.write_report(payload, {"results": []}, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Vest"]
    headers = [c.value for c in ws[1]]
    assert headers[-2:] == ["Status", "Note"]
    assert len(headers) == len(xlsxCreator.TEMPLATE_HEADERS) + 2

    # One line per variant, which is what the parser reads back. This payload
    # predates the template change and states no pairing, so the child SKUs are
    # recovered by combination_skus on the way out — see row_to_lines.
    row = next(r for r in payload if r["sku"] == "CP-V-200")
    lines = [r for r in ws.iter_rows(min_row=2, values_only=True)
             if r[headers.index("SKU")] == "CP-V-200"]
    assert len(lines) == len(row["colors"]) * len(row["sizes"])

    assert {line[headers.index("Product name")] for line in lines} == \
           {row["product_name"]}
    assert {line[headers.index("Color")] for line in lines} == set(row["colors"])
    assert {line[headers.index("Size")] for line in lines} == set(row["sizes"])
    assert {line[headers.index("Child SKU Varient")] for line in lines} == \
           set(row["combinations"])
