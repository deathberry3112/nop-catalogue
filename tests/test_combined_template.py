#!/usr/bin/env python3
"""One sheet carrying BOTH templates — phase 1 and phase 2 on the same lines.

    Product name | SKU | ... | Child SKU Varient | Save Product | Configurator
    Areas | Add new location | Location Name | Picario RefrenceId | Save |
    Picario Name | Thumbnail Name | Save | Save Product

Before this, a sheet like that was routed to the configurator reader alone —
looks_like_configurator matches on Location Name plus the Picario reference,
and both are there — which then reported all nine product columns as
unrecognised and created nothing.

TWO THINGS MAKE IT WORK, and every test here is about one of them.

THE BOUNDARY. Ownership of a column is positional: the leftmost phase-2 header
is where the configurator block starts, everything left of it is phase 1's, and
Product name / SKU / Status / Note are shared. That is what resolves the
duplicated "Save Product" header — the client's combined template carries it
twice, once per half — with no special case at all.

THE GRANULARITY. Phase 1 writes one line per variant and phase 2 one line per
area, so a product with seven sizes and four areas is seven lines of which four
also name a location. Neither reader may treat the other's blanks as its own:
lines 5-7 are sizes, not locations with no name.

The result is ONE row per product, because two rows would put the phases on
separate lines of the report and a re-uploaded report would no longer be the
combined template the user wrote. test_a_written_report_parses_back_unchanged
is that property.
"""

from __future__ import annotations

import openpyxl
import pytest

import xlsxCreator as xc
from app import configurator
from app.parser import parse_workbook

HEADERS = [
    "Product name", "SKU", "Attribute Specification Value",
    "Product Attributes Type", "Color", "Color Swatch",
    "Product Attributes Type", "Size",
    "Generate Combinations", "Allow Out Of Stock", "Child SKU Varient",
    "Save Product",
    "Configurator Areas", "Add new location", "Location Name",
    "Picario RefrenceId", "Save ", "Picario Name", "Thumbnail Name", "Save ",
    "Save Product",
]


def _line(size, child, *, colour="Lime/Black", location="", reference="",
          picario="", thumbnail=""):
    """One line of the combined sheet: a variant, optionally also an area."""
    return ["G347LA - M", "G347LA - M", "ZingHighVis",
            "Color", colour, colour, "Size", size,
            "Yes", "Yes", child, "Yes",
            "Select", "Yes" if reference else "", location, reference,
            "Yes" if reference else "", picario, thumbnail, "Yes", "Yes"]


#: Two locations of two areas each, across the first four of seven size lines.
SHEET = [
    _line("M", "G347LA6", location="Front", reference="G347LA-Front.pfs",
          picario="L-chest-upper-right", thumbnail="L-chest-upper-right.png"),
    _line("L", "G347LA5", location="Front",
          picario="K-chest-upper-left", thumbnail="K-chest-upper-left.png"),
    _line("XL", "G347LA0", location="Back", reference="G347LA-Back.pfs",
          picario="A-back-upper", thumbnail="A-back-upper.png"),
    _line("2XL", "G347LA1", location="Back",
          picario="B-back-lower", thumbnail="B-back-lower.png"),
    _line("3XL", "G347LA2"),
    _line("4XL", "G347LA3"),
    _line("5XL", "G347LA4"),
]


def _workbook(tmp_path, lines, headers=None, name="combined.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers or HEADERS)
    for line in lines:
        ws.append(line)
    path = tmp_path / name
    wb.save(path)
    return path


def _one(path):
    result = parse_workbook(path)
    assert result.errors == [], result.errors
    assert len(result.rows) == 1, [r["sku"] for r in result.rows]
    return result.rows[0]


# ---------------------------------------------------------------------------
# both halves, one row
# ---------------------------------------------------------------------------

def test_a_combined_sheet_yields_one_row_carrying_both_phases(tmp_path):
    row = _one(_workbook(tmp_path, SHEET))
    assert row["errors"] == []
    assert [(a["name"], a["values"]) for a in row["attributes"]] == [
        ("Color", ["Lime/Black"]),
        ("Size", ["M", "L", "XL", "2XL", "3XL", "4XL", "5XL"])]
    assert [(l["name"], l["reference_id"], len(l["areas"]))
            for l in row["locations"]] == [
        ("Front", "G347LA-Front.pfs", 2), ("Back", "G347LA-Back.pfs", 2)]


def test_the_product_columns_are_not_reported_as_unrecognised(tmp_path):
    """The bug this whole feature starts from: the phase-2 reader owned the
    sheet, so every phase-1 column was an unknown header and nothing ran."""
    row = _one(_workbook(tmp_path, SHEET))
    assert not any("unrecognised" in e for e in row["errors"])


def test_the_configurator_columns_are_not_reported_as_unrecognised(tmp_path):
    row = _one(_workbook(tmp_path, SHEET))
    assert not any("Location Name" in e for e in row["errors"])


def test_it_stays_a_product_row_so_the_runner_creates_before_configuring(tmp_path):
    """One row, and it takes the product path. The configurator tab lives on a
    product, so the phases cannot be two rows whose order is up to the batch."""
    row = _one(_workbook(tmp_path, SHEET))
    assert row.get("kind") != "configurator"
    assert row["locations"]


# ---------------------------------------------------------------------------
# the boundary
# ---------------------------------------------------------------------------

def test_the_boundary_is_the_leftmost_configurator_column():
    assert configurator.block_starts_at(tuple(HEADERS)) == 12
    assert HEADERS[12] == "Configurator Areas"


def test_a_product_only_sheet_has_no_boundary():
    assert configurator.block_starts_at(tuple(HEADERS[:12])) is None


def test_the_duplicated_save_product_column_is_split_by_position(tmp_path):
    """The client's template spells "Save Product" twice, once per half. Header
    matching alone takes the last, so phase 1 read phase 2's column — and the
    lines with no configurator data read as Save Product: No, failing the
    product with "the lines disagree about Save Product"."""
    lines = [list(line) for line in SHEET]
    for line in lines:
        line[11] = "Yes"        # phase 1's Save Product
        line[20] = ""           # phase 2's, blank on the size-only lines
    row = _one(_workbook(tmp_path, lines))
    assert row["errors"] == []
    assert row["save_product"] is True


def test_a_phase_two_sheet_is_still_read_as_one(tmp_path):
    """Product name and SKU are on both templates, so their presence cannot be
    what marks a sheet combined — a product-only column to the left of the
    boundary is."""
    headers = ["Product name", "SKU", "Configurator Areas", "Add new location",
               "Location Name", "Picario RefrenceId", "Save ", "Picario Name",
               "Thumbnail Name", "Save ", "Save Product"]
    lines = [["G347LA", "G347LA", "Select", "Yes", "Front", "G347LA-Front.pfs",
              "Yes", "L-chest-upper-right", "L-chest-upper-right.png", "Yes", "Yes"]]
    row = _one(_workbook(tmp_path, lines, headers=headers))
    assert row["kind"] == "configurator"


# ---------------------------------------------------------------------------
# the granularity
# ---------------------------------------------------------------------------

def test_a_size_only_line_is_not_a_location_with_no_name(tmp_path):
    """Lines 5-7 carry a SKU and a size and no configurator data. The phase-2
    reader's "is this line blank?" test asks whether any mapped cell is filled,
    and the SKU is one — so without the location filter these read as three
    nameless locations, which is exactly what the user's sheet reported."""
    row = _one(_workbook(tmp_path, SHEET))
    assert not any("no Location Name" in e for e in row["errors"])
    assert len(row["locations"]) == 2


def test_navigation_cells_alone_do_not_make_a_location(tmp_path):
    """'Configurator Areas: Select' is an instruction to a human, and the sheet
    repeats it down every line including the ones with no location."""
    row = _one(_workbook(tmp_path, SHEET))
    assert sum(len(l["areas"]) for l in row["locations"]) == 4


def test_more_areas_than_variants_still_reads_every_area(tmp_path):
    """The two halves are independent, so neither caps the other."""
    lines = [
        _line("M", "G347LA6", location="Front", reference="F.pfs",
              picario="one", thumbnail="one.png"),
        _line("", "", location="Front", picario="two", thumbnail="two.png"),
        _line("", "", location="Front", picario="three", thumbnail="three.png"),
    ]
    row = _one(_workbook(tmp_path, lines))
    assert [a["picario_name"] for a in row["locations"][0]["areas"]] == [
        "one", "two", "three"]


# ---------------------------------------------------------------------------
# the configurator's own validation still applies
# ---------------------------------------------------------------------------

def test_a_location_with_no_reference_still_fails_the_product(tmp_path):
    """The locations moved onto the product row; their errors have to move with
    them, or a broken location passes on a sheet that also creates the product."""
    lines = [_line("M", "G347LA6", location="Front",
                   picario="one", thumbnail="one.png")]
    result = parse_workbook(_workbook(tmp_path, lines))
    row = result.rows[0]
    assert any("no Picario RefrenceId" in e for e in row["errors"])


def test_a_gap_in_the_variant_grid_still_fails_the_product(tmp_path):
    """Phase 1's checks are untouched by the phase-2 half sharing the lines. Two
    colours over two sizes needs four lines; the fourth is missing here."""
    lines = [
        _line("M", "A1", colour="Lime", location="Front", reference="F.pfs",
              picario="one", thumbnail="one.png"),
        _line("L", "A2", colour="Lime"),
        _line("M", "A3", colour="Orange"),
    ]
    result = parse_workbook(_workbook(tmp_path, lines))
    assert any("no child SKU" in e for e in result.rows[0]["errors"])


def test_a_line_with_a_sku_and_no_product_data_is_still_one_row(tmp_path):
    """A combined sheet says "just configure this one" by leaving the product
    half blank. It stays a PRODUCT row — so the ownership guard still applies to
    it, and a SKU this automation did not create is left alone rather than
    configured, which is the phase-2 exception and not something a combined
    sheet should acquire by accident."""
    lines = list(SHEET) + [
        ["", "OTHER-1", "", "", "", "", "", "", "", "", "", "",
         "Select", "Yes", "Front", "O.pfs", "Yes", "spot", "spot.png", "Yes", "Yes"]]
    rows = parse_workbook(_workbook(tmp_path, lines)).rows

    other = next(r for r in rows if r["sku"] == "OTHER-1")
    assert other.get("kind") != "configurator"
    assert other["errors"] == []
    assert [l["name"] for l in other["locations"]] == ["Front"]
    assert not any(a["values"] for a in other["attributes"])


# ---------------------------------------------------------------------------
# the report round trip
# ---------------------------------------------------------------------------

def _shape(rows):
    return [{
        "sku": r["sku"],
        "attributes": [(a["name"], tuple(a["values"])) for a in r.get("attributes", [])],
        "combinations": tuple(r.get("combinations") or []),
        "variants": [(tuple(sorted(v["values"].items())), v["sku"])
                     for v in r.get("variants") or []],
        "locations": [(l["name"], l["reference_id"],
                       tuple((a["picario_name"], a["thumbnail"]) for a in l["areas"]))
                      for l in r.get("locations") or []],
        "errors": r["errors"],
    } for r in rows]


def test_a_written_report_parses_back_unchanged(tmp_path):
    """The reason a combined sheet is ONE row. The report is meant to be
    re-uploaded to finish a batch, so it has to come back out as the same
    template it went in as — both halves, on the same lines."""
    source = _workbook(tmp_path, SHEET)
    before = parse_workbook(source).rows

    report = tmp_path / "report.xlsx"
    xc.write_report(before, None, report)

    assert _shape(parse_workbook(report).rows) == _shape(before)


def test_the_report_keeps_the_combined_columns(tmp_path):
    source = _workbook(tmp_path, SHEET)
    report = tmp_path / "report.xlsx"
    xc.write_report(parse_workbook(source).rows, None, report)

    written = [c.value for c in openpyxl.load_workbook(report).worksheets[0][1]]
    assert written == HEADERS[:12] + HEADERS[12:] + ["Status", "Note"]


def test_a_products_own_attributes_name_the_columns(tmp_path):
    """The same rule the product template follows: a sheet is written back with
    the attributes it declared, not with Color and Size forced onto it."""
    headers = [h.replace("Color", "Sign Material").replace("Size", "Sign Size")
               for h in HEADERS]
    lines = [[c for c in line] for line in SHEET]
    for line in lines:
        line[3], line[6] = "Sign Material", "Sign Size"
    report = tmp_path / "report.xlsx"
    xc.write_report(parse_workbook(_workbook(tmp_path, lines, headers=headers)).rows,
                    None, report)

    written = [c.value for c in openpyxl.load_workbook(report).worksheets[0][1]]
    assert written[:8] == ["Product name", "SKU", "Attribute Specification Value",
                           "Product Attributes Type", "Sign Material",
                           "Sign Material Swatch", "Product Attributes Type",
                           "Sign Size"]
    assert "Location Name" in written


def test_a_line_carrying_only_a_status_is_not_a_variant(tmp_path):
    """A report's phase-2-only line has a Status and a Note and no product data.
    Counting that as a variant answers one failed line with two failed rows."""
    lines = list(SHEET) + [
        ["", "", "", "", "", "", "", "", "", "", "", "",
         "Select", "Yes", "Front", "", "Yes", "spot", "spot.png", "Yes", "Yes",
         "Failed", "SKU is required"]]
    rows = parse_workbook(_workbook(tmp_path, lines, headers=HEADERS + ["Status", "Note"])).rows
    assert [r["sku"] for r in rows] == ["G347LA - M", ""]
    assert rows[1]["kind"] == "configurator"
