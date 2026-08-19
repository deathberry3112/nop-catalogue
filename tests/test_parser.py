#!/usr/bin/env python3
"""The workbook parser — and specifically what it refuses to do.

One sheet line is one VARIANT: a colour, a size and the child SKU that pairs
them. Lines sharing a SKU are one product. Most of what is worth testing here
is the assembly — that lines become the right products, and that a sheet which
cannot mean what it appears to mean is refused rather than read.

The round-trip tests matter because parser.py and xlsxCreator.py are mirrors of
each other: one reads the 12 template columns, the other writes them. A change
to one that is not made to the other shows up here rather than in production.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

import xlsxCreator
from app.parser import parse_workbook

HEADERS = xlsxCreator.TEMPLATE_HEADERS


def _workbook(rows: list[list], headers: list[str] | None = None,
              sheet: str = "Products") -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers or HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _line(sku="SKU-1", color="Red", size="", combo="SKU-1 RD", *,
          name=None, swatch=None, spec="Zing", gen="Yes", oos="No",
          save="Yes", color_attr="Color", size_attr="Size") -> list:
    """One variant line: one colour, one size, one child SKU."""
    return [name if name is not None else "Product One", sku, spec, color_attr,
            color, swatch if swatch is not None else color, size_attr, size,
            gen, oos, combo, save]


def _product(sku="SKU-1", colors=("Red", "Blue"), sizes=("",), **kw) -> list[list]:
    """Every line of one product — the full colour x size grid."""
    return [_line(sku=sku, color=c, size=s, combo=f"{sku} {c[:2].upper()}{s}", **kw)
            for c in colors for s in sizes]


# --- assembling variant lines into products ---------------------------------

def test_lines_sharing_a_sku_become_one_product():
    result = parse_workbook(_workbook(_product()))
    assert len(result.rows) == 1
    row = result.rows[0]
    assert row["sku"] == "SKU-1"
    assert row["colors"] == ["Red", "Blue"]
    assert row["combinations"] == ["SKU-1 RE", "SKU-1 BL"]
    assert row["generate_combinations"] is True
    assert row["errors"] == []


def test_the_sheet_states_which_child_sku_pairs_with_which_combination():
    """The whole point of the template change. Nothing is inferred from the
    child SKU text any more — the line says what it belongs to."""
    result = parse_workbook(_workbook(_product(sizes=("S", "M"))))
    assert result.rows[0]["variants"][:2] == [
        {"color": "Red", "size": "S", "sku": "SKU-1 RES", "source_row": 2},
        {"color": "Red", "size": "M", "sku": "SKU-1 REM", "source_row": 3},
    ]


def test_the_grouping_key_is_the_sku_not_the_product_name():
    """Two products may share a display name and still be two products. The SKU
    is what this automation looks a product up by, so it decides identity."""
    lines = [_line(sku="SKU-1", color="Red", combo="SKU-1 RD", name="Shared"),
             _line(sku="SKU-2", color="Blue", combo="SKU-2 BE", name="Shared")]
    result = parse_workbook(_workbook(lines))
    assert [r["sku"] for r in result.rows] == ["SKU-1", "SKU-2"]
    assert all(r["product_name"] == "Shared" for r in result.rows)
    assert all(r["errors"] == [] for r in result.rows)


def test_a_products_lines_need_not_be_adjacent():
    lines = [_line(sku="A", color="Red", combo="A RD"),
             _line(sku="B", color="Red", combo="B RD"),
             _line(sku="A", color="Blue", combo="A BE")]
    result = parse_workbook(_workbook(lines))
    row = next(r for r in result.rows if r["sku"] == "A")
    assert row["colors"] == ["Red", "Blue"]
    assert row["errors"] == []


def test_an_empty_size_column_means_a_colour_only_product():
    result = parse_workbook(_workbook(_product(sizes=("",))))
    row = result.rows[0]
    assert row["sizes"] == []
    assert [v["size"] for v in row["variants"]] == [None, None]


def test_each_colour_keeps_its_own_swatch():
    lines = [_line(sku="S", color="Lime", swatch="Hi-Vis Lime", combo="S LE"),
             _line(sku="S", color="Orange", swatch="Hi-Vis Orange", combo="S OE")]
    row = parse_workbook(_workbook(lines)).rows[0]
    assert row["colors"] == ["Lime", "Orange"]
    assert row["swatches"] == ["Hi-Vis Lime", "Hi-Vis Orange"]
    assert row["errors"] == []


def test_every_sheet_is_read_and_attributed():
    wb = openpyxl.Workbook()
    for name in ("HardHat", "Vest"):
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for line in _product(sku=f"{name}-1"):
            ws.append(line)
    wb.remove(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    result = parse_workbook(buf)
    assert result.sheets == ["HardHat", "Vest"]
    assert {r["source_sheet"] for r in result.rows} == {"HardHat", "Vest"}


def test_blank_lines_are_ignored():
    result = parse_workbook(_workbook(_product() + [[None] * 12, [""] * 12]))
    assert len(result.rows) == 1
    assert result.rows[0]["source_rows"] == [2, 3]


# --- the refusals: each of these would otherwise be a silent wrong result ----

def test_a_sheet_in_the_old_format_is_refused_not_misread():
    """The two templates use identical columns and disagree only about what a
    line MEANS, so a legacy sheet read as variants would come out as a product
    per line — colour lists where a colour goes, child SKUs paired with
    nothing — and pass every other check on the way to the runner."""
    legacy = _line(color="Red, Blue", combo="SKU-1 RD\nSKU-1 BE")
    result = parse_workbook(_workbook([legacy]))
    assert result.rows == []
    assert "OLD template" in result.errors[0]
    assert "2 colours" in result.errors[0] and "2 child SKUs" in result.errors[0]


def test_the_refusal_names_the_row_so_it_can_be_found():
    result = parse_workbook(_workbook(_product() + [_line(sku="X", color="A, B",
                                                          combo="X A")]))
    assert "row 4" in result.errors[0] and "X" in result.errors[0]


def test_an_unrecognised_column_blocks_the_row():
    """The important one. An unknown column is very likely an attribute this
    pipeline cannot model, and creating the product without it yields a wrong
    catalogue and a clean report — the worst outcome available."""
    lines = [line + ["Leather"] for line in _product()]
    result = parse_workbook(_workbook(lines, headers=HEADERS + ["Material"]))
    assert result.rows[0]["errors"], "an unknown column must block the row"
    assert "Material" in result.rows[0]["errors"][0]


def test_a_missing_sku_is_an_error():
    result = parse_workbook(_workbook([_line(sku="")]))
    assert any("SKU is required" in e for e in result.rows[0]["errors"])


def test_lines_with_no_sku_are_reported_separately():
    """One error per mistake. Two blank-SKU lines are probably two unrelated
    slips, not one product."""
    result = parse_workbook(_workbook([_line(sku=""), _line(sku="")]))
    assert len(result.rows) == 2


def test_a_combination_with_no_child_sku_is_refused():
    """The check the old template could not make. It had counts and no pairing,
    so it could only say "8 children for 8 combinations"; this names the gap."""
    lines = _product(colors=("Red", "Blue"), sizes=("S", "M"))
    del lines[1]                                   # Red / M
    result = parse_workbook(_workbook(lines))
    error = next(e for e in result.rows[0]["errors"] if "no child SKU" in e)
    assert "Red / M" in error


def test_two_child_skus_for_one_combination_is_refused():
    lines = _product(colors=("Red",), sizes=("S",))
    lines.append(_line(color="Red", size="S", combo="SKU-1 OTHER"))
    result = parse_workbook(_workbook(lines))
    assert any("One combination takes one child SKU" in e
               for e in result.rows[0]["errors"])


def test_one_child_sku_used_twice_is_refused():
    lines = [_line(color="Red", combo="SKU-1 SAME"),
             _line(color="Blue", combo="SKU-1 SAME")]
    result = parse_workbook(_workbook(lines))
    assert any("more than one combination" in e for e in result.rows[0]["errors"])


def test_lines_disagreeing_about_the_product_are_refused():
    """These columns describe the product, so its lines have to agree. One of
    the two answers would otherwise be discarded, and nothing knows which."""
    lines = [_line(color="Red", combo="SKU-1 RD", name="Vest"),
             _line(color="Blue", combo="SKU-1 BE", name="Hard Hat")]
    result = parse_workbook(_workbook(lines))
    error = next(e for e in result.rows[0]["errors"] if "disagree" in e)
    assert "Product name" in error and "'Vest'" in error and "[3]" in error


def test_lines_disagreeing_about_a_flag_are_refused():
    lines = [_line(color="Red", combo="SKU-1 RD", save="Yes"),
             _line(color="Blue", combo="SKU-1 BE", save="No")]
    result = parse_workbook(_workbook(lines))
    assert any("Save Product" in e for e in result.rows[0]["errors"])


def test_one_colour_given_two_swatches_is_refused():
    lines = [_line(sku="S", color="Lime", swatch="Hi-Vis Lime", size="S",
                   combo="S LE S"),
             _line(sku="S", color="Lime", swatch="Pale Lime", size="M",
                   combo="S LE M")]
    result = parse_workbook(_workbook(lines))
    assert any("different swatches" in e for e in result.rows[0]["errors"])


def test_a_product_half_with_sizes_and_half_without_is_refused():
    lines = [_line(color="Red", size="S", combo="SKU-1 RD S"),
             _line(color="Blue", size="", combo="SKU-1 BE")]
    result = parse_workbook(_workbook(lines))
    assert any("sizes on every line or on none" in e
               for e in result.rows[0]["errors"])


def test_a_duplicate_sku_across_sheets_is_flagged_on_both():
    wb = openpyxl.Workbook()
    for name in ("A", "B"):
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for line in _product():
            ws.append(line)
    wb.remove(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    result = parse_workbook(buf)
    assert all(any("more than one sheet" in e for e in r["errors"])
               for r in result.rows)


def test_generate_without_colours_is_refused():
    result = parse_workbook(_workbook([_line(color="", combo="")]))
    assert any("no child SKUs" in e for e in result.rows[0]["errors"])


def test_a_child_sku_with_no_colour_is_refused():
    result = parse_workbook(_workbook([_line(color="", combo="SKU-1 RD")]))
    assert any("no colour" in e for e in result.rows[0]["errors"])


def test_a_clean_product_carries_no_duplicate_payload():
    assert parse_workbook(_workbook(_product())).rows[0]["duplicates"] == {}


def test_a_parent_sku_its_children_do_not_extend_is_not_nagged_about():
    """Dropped with the old template. The prefix was load-bearing then — it was
    what base_prefix stripped to recover the colour code — and carries nothing
    now that the sheet states the pairing. The client's own convention breaks
    it on every row: parent 'CP-280-HP241R - M', child 'CP-280-HP241R-01'."""
    lines = [_line(sku="CP-280 - M", color="White", combo="CP-280-01"),
             _line(sku="CP-280 - M", color="Red", combo="CP-280-15")]
    row = parse_workbook(_workbook(lines)).rows[0]
    assert row["errors"] == [] and row["warnings"] == []


# --- the attribute columns --------------------------------------------------

def test_a_size_column_repeating_the_colour_attribute_is_corrected():
    """Both attribute columns saying "Color" would make add_attribute hand back
    the colour attribute for the size pass, filing "OS" as a ninth colour."""
    lines = _product(sizes=("OS",), size_attr="Color")
    result = parse_workbook(_workbook(lines))
    assert result.rows[0]["size_attr"] == "Size"
    assert result.rows[0]["warnings"]

    # said once for the sheet, in the words that appear in the spreadsheet
    assert len(result.sheet_warnings) == 1
    note = result.sheet_warnings[0]
    assert "Product Attributes Type" in note
    assert "'Size'" in note


def test_a_sheet_with_only_one_attribute_column_is_told_to_add_the_second():
    headers = [h for h in HEADERS]
    del headers[6]                                 # drop the size attribute column
    lines = [line[:6] + line[7:] for line in _product(sizes=("OS",))]
    result = parse_workbook(_workbook(lines, headers=headers))
    assert result.rows[0]["size_attr"] == "Size"
    note = result.sheet_warnings[0]
    assert "only ONE 'Product Attributes Type'" in note
    assert "before the Size column" in note


def test_a_sheet_without_sizes_is_not_nagged_about_a_size_attribute():
    """A hard-hat sheet leaves Size empty. There is no size attribute to name,
    so there is nothing to fix and nothing to say."""
    headers = [h for h in HEADERS]
    del headers[6]
    lines = [line[:6] + line[7:] for line in _product(sizes=("",))]
    result = parse_workbook(_workbook(lines, headers=headers))
    assert result.sheet_warnings == []


def test_a_normal_size_attribute_is_left_alone():
    result = parse_workbook(_workbook(_product(sizes=("OS",))))
    assert result.rows[0]["size_attr"] == "Size"
    assert result.rows[0]["warnings"] == []
    assert result.sheet_warnings == []


def test_a_sheet_without_a_sku_column_is_reported_not_ignored():
    result = parse_workbook(_workbook([["a", "b"]], headers=["Product name", "Note"]))
    assert result.errors and "SKU" in result.errors[0]


# --- the round trip ---------------------------------------------------------

def test_a_written_report_parses_back(tmp_path):
    """xlsxCreator writes it, the parser reads it, and the fields survive.
    This is what lets a user re-upload a report to finish a partial batch."""
    original = {
        "sku": "CP-V-200", "product_name": "Vest", "spec_value": "ZingHighViz",
        "color_attr": "Color", "colors": ["Lime", "Orange"],
        "swatches": ["Lime", "Orange"], "size_attr": "Size",
        "sizes": ["S", "M"], "generate_combinations": True,
        "allow_out_of_stock": True, "save_product": True,
        "combinations": ["CP-V-200 LE S", "CP-V-200 LE M",
                         "CP-V-200 OE S", "CP-V-200 OE M"],
        "variants": [
            {"color": "Lime", "size": "S", "sku": "CP-V-200 LE S"},
            {"color": "Lime", "size": "M", "sku": "CP-V-200 LE M"},
            {"color": "Orange", "size": "S", "sku": "CP-V-200 OE S"},
            {"color": "Orange", "size": "M", "sku": "CP-V-200 OE M"},
        ],
        "source_sheet": "Vest", "errors": [],
    }
    path = tmp_path / "report.xlsx"
    xlsxCreator.write_report([original], None, path)

    parsed = parse_workbook(path).rows[0]
    for field in ("sku", "product_name", "spec_value", "colors", "sizes",
                  "combinations", "generate_combinations", "allow_out_of_stock"):
        assert parsed[field] == original[field], field
    assert parsed["source_sheet"] == "Vest"
    assert parsed["errors"] == []
    assert [(v["color"], v["size"], v["sku"]) for v in parsed["variants"]] == \
           [(v["color"], v["size"], v["sku"]) for v in original["variants"]]


def test_a_report_is_written_one_line_per_variant(tmp_path):
    """The mirror of the parser, and the reason a report can be re-uploaded: a
    report written in the old aggregated shape would come back as a workbook
    the parser now refuses."""
    row = {
        "sku": "S", "product_name": "P", "colors": ["Red", "Blue"],
        "sizes": ["S"], "combinations": ["S RD S", "S BE S"],
        "variants": [{"color": "Red", "size": "S", "sku": "S RD S"},
                     {"color": "Blue", "size": "S", "sku": "S BE S"}],
        "source_sheet": "Sheet", "errors": [],
    }
    path = tmp_path / "report.xlsx"
    counts = xlsxCreator.write_report([row], None, path)

    ws = openpyxl.load_workbook(path).active
    assert ws.max_row == 3                          # header + 2 variants
    assert [c.value for c in ws[2]][4] == "Red"
    assert [c.value for c in ws[3]][4] == "Blue"
    assert counts == {"Pending": 1}, "counted per product, not per line"


def test_a_carried_status_survives_so_a_rerun_can_skip_the_row():
    lines = [line + ["Added", ""] for line in _product()]
    result = parse_workbook(_workbook(lines, headers=HEADERS + ["Status", "Note"]))
    assert result.rows[0]["status"] == "Added"
    import main
    assert main._already_done(result.rows[0])


def test_a_product_not_generating_combinations_is_not_held_to_the_grid():
    """The old parser only checked the grid when Generate Combinations was set,
    and the runner still returns before it builds one. Blocking a row over a
    gap in a grid it will never generate would refuse work that runs fine."""
    lines = [_line(color="Red", size="S", combo="SKU-1 RD S", gen="No"),
             _line(color="Blue", size="M", combo="SKU-1 BE M", gen="No")]
    result = parse_workbook(_workbook(lines))
    assert result.rows[0]["errors"] == []
    assert result.valid_rows == 1


def test_a_product_generating_combinations_is():
    lines = [_line(color="Red", size="S", combo="SKU-1 RD S"),
             _line(color="Blue", size="M", combo="SKU-1 BE M")]
    result = parse_workbook(_workbook(lines))
    assert any("no child SKU" in e for e in result.rows[0]["errors"])
